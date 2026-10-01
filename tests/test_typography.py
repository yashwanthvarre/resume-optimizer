"""Downloads: EB Garamond everywhere, justified body text, bold keywords, ATS-safe .docx (AI mocked)."""
import os, sys, re, tempfile, zipfile
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
os.environ["ANTHROPIC_API_KEY"] = "test"
_home = tempfile.mkdtemp(); os.environ["HOME"] = os.environ["USERPROFILE"] = _home
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from fastapi.testclient import TestClient
import app as appmod
from ro import analyzer, resume_io as R
sys.path.insert(0, str(Path(__file__).parent))
from test_e2e_data import JD, CH, COVER

analyzer._call_tool = lambda s, u, tool, mt: {"record_jd": JD, "record_changes": CH, "record_cover_letter": COVER}[tool["name"]]
OUT = Path(tempfile.mkdtemp())
SRC = (Path(__file__).parent / "sample_resume.docx").resolve()
TERMS = ["Python", "REST APIs", "RESTful APIs", "AWS", "CI/CD"]
W_R = qn("w:r")


def fonts_ok(doc, font="EB Garamond"):
    """Every run and the defaults name the font in all four slots, with no theme font to override it."""
    bad = []
    for r in doc.element.body.iter(W_R):
        rf = r.rPr.rFonts if r.rPr is not None else None
        if rf is None or any(rf.get(qn(f"w:{k}")) != font for k in ("ascii", "hAnsi", "cs", "eastAsia")) \
                or any(rf.get(qn(f"w:{k}")) for k in ("asciiTheme", "hAnsiTheme")):
            bad.append("".join(t.text or "" for t in r.iter(qn("w:t"))))
    defaults = doc.styles.element.find(qn("w:docDefaults")).find(qn("w:rPrDefault")).find(qn("w:rPr")).rFonts
    assert defaults.get(qn("w:ascii")) == font
    assert doc.styles["Normal"].font.name == font
    return bad


def bold_words(p):
    return [r.text for r in p.runs if r.bold and r.text.strip()]


def export(edits=None, paras=None, **kw):
    d = R.parse_resume(SRC)
    dest = R.output_path(OUT, "Resume")
    R.export_docx(SRC, "docx", paras or d["paragraphs"], edits or {}, dest, **kw)
    return d, Document(str(dest))


def test_font_everywhere():
    _, doc = export(keywords=TERMS)
    assert fonts_ok(doc) == []
    _, kept = export(keywords=TERMS, font=R.KEEP_FONT)
    assert not any(r.rPr is not None and r.rPr.rFonts is not None and r.rPr.rFonts.get(qn("w:ascii")) == "EB Garamond"
                   for r in kept.element.body.iter(W_R)), "keep = the resume's own font"


def test_justify_body_only():
    long = "Software engineer with 4 years of experience building web backends in Python, owning services end to end."
    d, doc = export({"p3": long, "p6": "Built services in Python and Flask that handled customer orders across three regions with no downtime."})
    by = {p.text: p.paragraph_format.alignment for p in doc.paragraphs}
    assert by[long] == WD_ALIGN_PARAGRAPH.JUSTIFY
    assert by["Built services in Python and Flask that handled customer orders across three regions with no downtime."] == WD_ALIGN_PARAGRAPH.JUSTIFY
    assert by["ALEX SAMPLE"] != WD_ALIGN_PARAGRAPH.JUSTIFY                       # name
    assert by["alex@example.com | (555) 010-2000 | linkedin.com/in/alexsample"] != WD_ALIGN_PARAGRAPH.JUSTIFY
    assert by["SUMMARY"] != WD_ALIGN_PARAGRAPH.JUSTIFY                            # heading
    assert by["Acme Corp\tSoftware Engineer\tJan 2022 – Present"] != WD_ALIGN_PARAGRAPH.JUSTIFY  # tab row
    assert by["Python, Flask, SQL, AWS, Git, Docker"] != WD_ALIGN_PARAGRAPH.JUSTIFY  # one-liner
    st = doc.settings.element
    assert st.find(qn("w:autoHyphenation")) is not None and st.find(qn("w:doNotHyphenateCaps")) is not None
    assert not R.justify_ok("text", "Remote · New York, NY · 2019 – 2021 and a long enough line to wrap on the page here")
    assert not R.justify_ok("text", "First line\nsecond line of a manually broken paragraph that is long enough to wrap")


def test_bold_keywords_only():
    d, doc = export({"p7": "Owned AWS deployments through CI/CD and fixed production bugs, cutting incidents by 30%."}, keywords=TERMS)
    texts = [p.text for p in doc.paragraphs]
    orig = [R.applied_text([p], {}) for p in d["paragraphs"]]
    ps = {p.text: p for p in doc.paragraphs}
    p7 = ps["Owned AWS deployments through CI/CD and fixed production bugs, cutting incidents by 30%."]
    assert bold_words(p7) == ["AWS", "CI/CD"], bold_words(p7)
    assert bold_words(ps["Python, Flask, SQL, AWS, Git, Docker"]) == ["Python", "AWS"]
    assert bold_words(ps["ALEX SAMPLE"]) == ["ALEX SAMPLE"]         # the name was bold already: untouched
    assert bold_words(ps["SUMMARY"]) in ([], ["SUMMARY"])            # headings get nothing new
    assert "Acme Corp" in bold_words(ps["Acme Corp\tSoftware Engineer\tJan 2022 – Present"])  # existing bold kept
    assert texts[:len(orig)] != [] and "ALEX SAMPLE" == texts[0]
    _, plain = export(keywords=TERMS, bold=False)
    assert bold_words(next(p for p in plain.paragraphs if p.text.startswith("Python, Flask"))) == []


def test_bold_split_keeps_format_and_hyperlinks():
    doc = Document()
    doc.add_paragraph("Name")
    doc.add_paragraph("EXPERIENCE")
    p = doc.add_paragraph()
    it = p.add_run("Shipped Python services and ")
    it.italic = True
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), p.part.relate_to("https://example.com", "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True))
    lr = OxmlElement("w:r"); lt = OxmlElement("w:t"); lt.text = "AWS tooling"; lr.append(lt); link.append(lr)
    p._p.append(link)
    p.add_run(" for teams.")
    paras = [{"id": "p0", "text": "Name", "kind": "text"}, {"id": "p1", "text": "EXPERIENCE", "kind": "heading"},
             {"id": "p2", "text": p.text, "kind": "text"}]
    before = "".join(t.text for t in p._p.iter(qn("w:t")))
    R.finish_resume(doc, paras, {"p2": p._p}, ["Python", "AWS"])
    after = "".join(t.text for t in p._p.iter(qn("w:t")))
    assert before == after == "Shipped Python services and AWS tooling for teams."
    runs = [(("".join(t.text for t in r.iter(qn("w:t")))), r.rPr is not None and r.rPr.b is not None,
             r.rPr is not None and r.rPr.i is not None, r.getparent().tag == qn("w:hyperlink")) for r in p._p.iter(W_R)]
    assert ("Python", True, True, False) in runs, runs          # bold, still italic
    assert ("Shipped ", False, True, False) in runs
    assert ("AWS", True, False, True) in runs, runs             # still inside the hyperlink
    assert (" tooling", False, False, True) in runs
    assert len(p._p.findall(qn("w:hyperlink"))) == 1


ATS_TAGS = ("<w:txbxContent", "<w:drawing", "<w:pict", "<w:tbl>", "<mc:AlternateContent")


def _layout(path):
    with zipfile.ZipFile(path) as z:
        xml = z.read("word/document.xml").decode()
        hf = [n for n in z.namelist() if re.match(r"word/(header|footer)\d*\.xml", n)]
        hf_text = "".join(re.sub(r"<[^>]+>", "", z.read(n).decode()) for n in hf).strip()
    return {t: xml.count(t) for t in ATS_TAGS}, hf_text


def ats_safe(path, source=None):
    """A rebuilt file is plain paragraphs only: no text boxes, shapes, images or tables, and nothing in the
    header/footer. A .docx edited in place keeps the candidate's own layout but must never add any of those."""
    got, hf_text = _layout(path)
    base, base_hf = _layout(source) if source and source.suffix == ".docx" else ({t: 0 for t in ATS_TAGS}, "")
    assert got == base, (path.name, got, base)
    assert hf_text == base_hf, hf_text


def test_ats_round_trip_docx_and_rebuilt():
    c = TestClient(appmod.app)
    c.post("/api/settings", json={"file_format": "docx"})   # these checks read the .docx; test_pdf.py covers PDF
    for src in (SRC, None):
        if src is None:   # PDF/TXT path: a clean rebuilt .docx
            txt = OUT / "resume.txt"
            txt.write_text("ALEX SAMPLE\nalex@example.com | (555) 010-2000\nSUMMARY\nSoftware engineer building Python services on AWS "
                           "for four years, owning delivery from design to production support.\nEXPERIENCE\n"
                           "• Built REST APIs in Python for customer orders and payments across two regions.\nSKILLS\nPython, AWS, SQL\n")
            src = txt
        sid = c.post("/api/resume/load", json={"path": str(src)}).json()["session_id"]
        res = c.post("/api/analyze", json={"session_id": sid, "jd_text": "x" * 200}).json()
        paras = c.post("/api/resume/load", json={"path": str(src)}).json()
        accepted = [{"target_id": x["target_id"], "new_text": x["new_text"]} for x in res["changes"]
                    if x["target_id"] != "p5" and src.suffix == ".docx"]   # canned ids only fit the sample .docx
        out = c.post("/api/export", json={"session_id": sid, "accepted": accepted, "output_dir": str(OUT)}).json()
        path = Path(out["path"])
        ats_safe(path, src)
        doc = Document(str(path))
        assert fonts_ok(doc) == []
        got = [t.strip() for t in map(R._p_text, R._docx_paragraphs(doc)) if t.strip()]   # incl. table cells
        edits = {a["target_id"]: a["new_text"] for a in accepted}
        want = [R.BULLET_RE.sub("", edits.get(p["id"], p["text"])).strip() if src.suffix != ".docx" else edits.get(p["id"], p["text"]).strip()
                for p in appmod.SESSIONS[sid]["paragraphs"] if p["text"].strip()]
        assert got == want, (got, want)
        assert got[0] == "ALEX SAMPLE" and "@" in got[1]            # name + contact first, in the body
        kw_para = next(p for p in doc.paragraphs if p.text.strip().startswith("Python"))
        assert any(r.bold for r in kw_para.runs if r.text.strip() == "Python")


def test_settings_and_letter():
    c = TestClient(appmod.app)
    c.post("/api/settings", json={"file_format": "docx"})   # these checks read the .docx; test_pdf.py covers PDF
    cfg = c.post("/api/settings", json={"doc_font": "keep", "bold_keywords": False}).json()
    assert cfg["doc_font"] == "keep" and cfg["bold_keywords"] is False
    sid = c.post("/api/resume/load", json={"path": str(SRC)}).json()["session_id"]
    c.post("/api/analyze", json={"session_id": sid, "jd_text": "x" * 200})
    out = c.post("/api/export", json={"session_id": sid, "accepted": [], "output_dir": str(OUT)}).json()
    doc = Document(out["path"])
    assert not any(r.bold for p in doc.paragraphs if p.text.startswith("Python, Flask") for r in p.runs)
    cfg = c.post("/api/settings", json={"doc_font": "EB Garamond", "bold_keywords": True}).json()
    assert cfg["doc_font"] == "EB Garamond" and cfg["bold_keywords"] is True
    letter = {"greeting": "Dear Ms. Rivera,", "closing": "Sincerely,", "signature": "Alex Sample",
              "paragraphs": ["I build Python services on AWS and would like to bring that to Globex's backend team, "
                             "where REST APIs carry most of the product."]}
    for flag in (False, True):
        out = c.post("/api/export", json={"session_id": sid, "accepted": [], "output_dir": str(OUT), "include_resume": False,
                                          "cover_letter": letter, "bold_letter_keywords": flag}).json()
        doc = Document(out["cover"]["path"])
        assert fonts_ok(doc) == []
        body = next(p for p in doc.paragraphs if p.text.startswith("I build"))
        assert body.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
        greet = next(p for p in doc.paragraphs if p.text.startswith("Dear"))
        assert greet.paragraph_format.alignment != WD_ALIGN_PARAGRAPH.JUSTIFY
        assert bold_words(body) == (["Python", "AWS", "REST APIs"] if flag else [])


if __name__ == "__main__":
    for name, fn in list(globals().items()):
        if name.startswith("test_"):
            fn(); print("ok", name)
    print("ALL TYPOGRAPHY TESTS PASSED")
