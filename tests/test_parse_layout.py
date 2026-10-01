"""Resumes converted from PDF often store a section title before the name (e.g. "SUMMARY", then the name as
Heading 1 after a column break). The name/contact block must still come first, and export must still edit
the right paragraphs."""
import sys, tempfile
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from ro import resume_io, analyzer

d = Path(tempfile.mkdtemp())
doc = Document()
for _ in range(2):
    doc.add_paragraph("")
doc.add_paragraph("SUMMARY")
name = doc.add_paragraph(style="Heading 1")
br = OxmlElement("w:br"); br.set(qn("w:type"), "column")
r = name.add_run(); r._r.append(br)
name.add_run("Jane Doe")
doc.add_paragraph("jane@example.com | 555-010-2000 | LinkedIn")
doc.add_paragraph("")
doc.add_paragraph("Frontend engineer with 5 years of React experience.")
doc.add_paragraph("EXPERIENCE")
doc.add_paragraph("ACME CORP\tJAN 2022 - PRESENT")
doc.add_paragraph("Senior Engineer")
doc.add_paragraph("• Built a design system in React.")
src = d / "converted.docx"; doc.save(str(src))

paras = resume_io.parse_resume(src)["paragraphs"]
shown = [(p["id"], p["kind"], p["section"], p["text"].strip()) for p in paras if p["text"].strip()]
print(shown)
assert shown[0][3] == "Jane Doe" and shown[0][1] == "text" and shown[0][2] == "Header", shown[0]
assert shown[1][3].startswith("jane@") and shown[1][2] == "Header"
assert shown[2] == ("p2", "heading", "Summary", "SUMMARY"), shown[2]
summary = next(p for p in paras if p["text"].startswith("Frontend engineer"))
assert summary["section"] == "Summary"
job = next(p for p in paras if p["text"].startswith("ACME"))
assert "\t" not in job["section"] and "·" in next(p for p in paras if p["text"].startswith("Senior"))["section"]
assert resume_io.contact_header(paras) == ["Jane Doe", "jane@example.com | 555-010-2000 | LinkedIn"]
print("ok: name/contact first, sections right")

# ids still point at the right paragraphs in the file
dest = d / "out.docx"
resume_io.export_docx(src, "docx", paras, {summary["id"]: "Frontend engineer with 5 years of React and TypeScript experience."}, dest)
texts = [p.text for p in Document(str(dest)).paragraphs]
assert "Frontend engineer with 5 years of React and TypeScript experience." in texts and texts[2] == "SUMMARY"
print("ok: export edits the right paragraph")

# change cards follow reading order, not file order
raw = {"changes": [{"target_id": summary["id"], "new_text": "Frontend engineer with 5+ years of React experience.", "type": "reword", "reason": "r"},
                   {"target_id": shown[0][0], "new_text": "Jane A. Doe", "type": "reword", "reason": "r"}]}
order = [c["target_id"] for c in analyzer.validate_changes(raw, paras)["changes"]]
assert order == [shown[0][0], summary["id"]], order
print("ok: reading order")

# a plain resume with the name first is unchanged
plain = resume_io.parse_resume(Path(__file__).parent / "sample_resume.docx")["paragraphs"]
assert [p["id"] for p in plain] == [f"p{i}" for i in range(len(plain))]
print("OK")
