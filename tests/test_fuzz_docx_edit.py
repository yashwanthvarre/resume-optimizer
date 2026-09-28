"""Fuzz test: word-level edits applied to multi-run paragraphs must yield exactly the new text."""
import random, sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from docx import Document
from ro.resume_io import _apply_text, _p_text

random.seed(1)
words = "built led designed Python AWS APIs the and team of scalable services reduced latency by 30% \t , .".split(" ")
fails = n = 0
for trial in range(5000):
    d = Document(); p = d.add_paragraph()
    for _ in range(random.randint(1, 6)):
        r = p.add_run(" ".join(random.choice(words) for _ in range(random.randint(0, 4))))
        r.bold = random.random() < .3
        if random.random() < .2: r.add_tab()
    old = _p_text(p._p)
    toks = old.split(" ")
    for _ in range(random.randint(1, 4)):
        op = random.choice("ids"); i = random.randrange(len(toks) + 1)
        w = random.choice(words)
        if op == "i": toks.insert(i, w)
        elif op == "d" and toks: toks.pop(min(i, len(toks) - 1))
        elif toks: toks[min(i, len(toks) - 1)] = w
    new = " ".join(toks)
    n += 1
    _apply_text(p._p, new)
    if _p_text(p._p) != new:
        fails += 1
        if fails < 4: print(repr(old), "->", repr(new), "got", repr(_p_text(p._p)))
print(f"fuzz: {fails} failures of {n}")
assert fails == 0
