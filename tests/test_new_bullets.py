"""Add it: a keyword no line can carry becomes a new bullet in the closest role, never a reason (AI mocked)."""
import os, re, sys, tempfile, zipfile
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
os.environ["ANTHROPIC_API_KEY"] = "test"
_home = tempfile.mkdtemp(); os.environ["HOME"] = os.environ["USERPROFILE"] = _home
from docx import Document
from fastapi.testclient import TestClient
import app as appmod
from ro import analyzer, resume_io as R
sys.path.insert(0, str(Path(__file__).parent))
from test_e2e_data import JD, CH, COVER, JD_TEXT

SRC = (Path(__file__).parent / "sample_resume.docx").resolve()
NEW = "Shared order and session state across the checkout views with contextAPI, removing prop drilling."
OVERRIDE = []
def fake(s, u, tool, mt):
    if tool["name"] == "record_keyword_decisions":
        return OVERRIDE.pop()
    return {"record_jd": JD, "record_changes": CH, "record_cover_letter": COVER}[tool["name"]]
analyzer._call_tool = fake
paras = R.parse_resume(SRC)["paragraphs"]


def test_merge():
    raw = {"changes": [], "new_bullets": [{"after_id": "p8", "text": "• " + NEW, "jd_keywords": ["contextAPI"]},
                                          {"after_id": "p5", "text": NEW, "jd_keywords": ["contextAPI"]},   # a header row
                                          {"after_id": "p99", "text": NEW, "jd_keywords": []}]}
    dropped = analyzer.merge_new_bullets(raw, paras, {})
    assert len(dropped) == 2, dropped
    [c] = raw["changes"]
    assert c["target_id"] == "p8" and c["type"] == "new_bullet"
    assert c["new_text"] == "Helped teammates with code reviews.\n" + NEW   # bullet symbol dropped, Word adds it
    # builds on Claude's own rewrite of the same bullet and on the candidate's accepted edit
    raw = {"changes": [{"target_id": "p8", "new_text": "Led code reviews.", "jd_keywords": ["code review"]}],
           "new_bullets": [{"after_id": "p8", "text": NEW, "jd_keywords": ["contextAPI"]}]}
    analyzer.merge_new_bullets(raw, paras, {})
    assert raw["changes"][0]["new_text"] == "Led code reviews.\n" + NEW
    assert raw["changes"][0]["jd_keywords"] == ["code review", "contextAPI"]
    raw = {"new_bullets": [{"after_id": "p8", "text": NEW, "jd_keywords": ["contextAPI"]}]}
    analyzer.merge_new_bullets(raw, paras, {"p8": "Mentored teammates through code reviews."})
    assert raw["changes"][0]["new_text"].startswith("Mentored teammates through code reviews.\n")


def test_endpoint_adds_a_bullet():
    c = TestClient(appmod.app)
    sid = c.post("/api/resume/load", json={"path": str(SRC)}).json()["session_id"]
    c.post("/api/analyze", json={"session_id": sid, "jd_text": JD_TEXT})
    OVERRIDE.append({"decisions": [{"term": "contextAPI", "evidence": "weak", "explanation": "New bullet at Acme."}],
                     "changes": [], "new_bullets": [{"after_id": "p8", "text": NEW, "jd_keywords": ["contextAPI"]}]})
    d = c.post("/api/justify_keywords", json={"session_id": sid, "items": [{"term": "contextAPI"}]})
    assert d.status_code == 200, d.text
    d = d.json()
    assert d["decisions"][0]["added"], d
    [ch] = d["changes"]
    assert ch["target_id"] == "p8" and ch["new_text"].endswith("\n" + NEW) and ch["source"] == "add"
    assert not any("longer" in w for w in ch["warnings"]), ch["warnings"]   # a new line isn't a longer bullet


def _texts(doc):
    return [p.text for p in doc.paragraphs if p.text.strip()]


def test_export_docx_and_clean():
    edits = {"p8": "Helped teammates with code reviews.\n" + NEW,
             "p5": "Acme Corp\tSoftware Engineer\tJan 2022 – Present"}   # not a bullet: never split
    out = Path(tempfile.mkdtemp())
    for kind in ("docx", "pdf"):
        dest = R.export_docx(SRC, kind, paras, edits, out / f"{kind}.docx", keywords=["contextAPI"])
        doc = Document(str(dest))
        t = _texts(doc)
        i = t.index("Helped teammates with code reviews.")
        assert t[i + 1] == NEW, (kind, t[i:i + 3])
        new, old = doc.paragraphs[[p.text for p in doc.paragraphs].index(NEW)], \
            doc.paragraphs[[p.text for p in doc.paragraphs].index(t[i])]
        assert new.style.name == old.style.name                      # same bullet formatting
        assert any(r.bold and r.text == "contextAPI" for r in new.runs)
    # the new paragraph doesn't reuse Word's paragraph ids
    body = zipfile.ZipFile(out / "docx.docx").read("word/document.xml").decode()
    ids = re.findall(r'w14:paraId="([0-9A-F]+)"', body)
    assert len(ids) == len(set(ids))


def test_split_keeps_existing_breaks():
    p = {"kind": "bullet", "text": "line one\nline two"}
    assert R.split_new_bullets(p, "line one\nline two\nnew one\n\nnew two") == ("line one\nline two", ["new one", "new two"])
    assert R.split_new_bullets({"kind": "text", "text": "a"}, "a\nb") == ("a\nb", [])


if __name__ == "__main__":
    for name, fn in list(globals().items()):
        if name.startswith("test_"):
            fn(); print("ok", name)
    print("ALL NEW-BULLET TESTS PASSED")
