"""Create a sample resume with mixed formatting for testing."""
from docx import Document
from docx.shared import Pt
from pathlib import Path
import sys

out = Path(sys.argv[1] if len(sys.argv) > 1 else "tests/sample_resume.docx")
d = Document()
p = d.add_paragraph(); r = p.add_run("ALEX SAMPLE"); r.bold = True; r.font.size = Pt(20)
d.add_paragraph("alex@example.com | (555) 010-2000 | linkedin.com/in/alexsample")
d.add_heading("SUMMARY", level=1)
d.add_paragraph("Software engineer with 4 years of experience building web backends in Python and working with cloud services.")
d.add_heading("EXPERIENCE", level=1)
p = d.add_paragraph(); r = p.add_run("Acme Corp"); r.bold = True; p.add_run("\tSoftware Engineer\tJan 2022 – Present")
d.add_paragraph("Built services in Python and Flask that handled customer orders.", style="List Bullet")
p = d.add_paragraph(style="List Bullet"); p.add_run("Worked on "); r = p.add_run("AWS"); r.italic = True; p.add_run(" deployments and fixed bugs, cutting incidents by 30%.")
d.add_paragraph("Helped teammates with code reviews.", style="List Bullet")
d.add_heading("SKILLS", level=1)
d.add_paragraph("Python, Flask, SQL, AWS, Git, Docker")
t = d.add_table(rows=1, cols=2); t.cell(0,0).text = "B.S. Computer Science"; t.cell(0,1).text = "2021"
d.save(out); print("wrote", out)
