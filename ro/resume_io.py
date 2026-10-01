"""Read resumes into editable paragraphs and write the optimized .docx.

DOCX: every paragraph (body, tables, text boxes) gets a stable id "p<n>". On export we
edit a *copy* of the original and only touch the characters that actually changed,
so fonts, bold/italic runs, bullets, tabs and layout survive.

PDF / TXT: layout can't be recovered, so we rebuild a clean ATS-friendly .docx.
"""
from __future__ import annotations

import copy
import difflib
import re
import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, Inches, RGBColor
from docx.oxml import OxmlElement

W_P, W_R, W_T = qn("w:p"), qn("w:r"), qn("w:t")
W_TAB, W_BR, W_CR = qn("w:tab"), qn("w:br"), qn("w:cr")
MC_FALLBACK = "{http://schemas.openxmlformats.org/markup-compatibility/2006}Fallback"
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"
BULLET_RE = re.compile(r"^\s*[•●▪■◦‣∙·\-–*\uf000-\uf0ff➢➤►▸✓✔]\s*(?=\S)")
PUA_BULLET = re.compile(r"^\s*[\uf000-\uf0ff]\s*")
DOC_FONT = "EB Garamond"
KEEP_FONT = "keep"
SEPARATOR_RE = re.compile(r"\t| {2,}\S| [·|] ")   # tab-aligned or " · " / " | " rows (dates, locations)


# --------------------------------------------------------------------------- DOCX internals
def _docx_paragraphs(doc) -> list:
    """All w:p elements in reading order, skipping duplicate VML fallbacks of text boxes."""
    out = []
    for p in doc.element.body.iter(W_P):
        if any(a.tag == MC_FALLBACK for a in p.iterancestors()):
            continue
        out.append(p)
    return out


def _segments(p) -> list[list]:
    """[(element, kind, text)] for runs that belong directly to paragraph p (incl. hyperlinks)."""
    segs = []
    for r in p.iter(W_R):
        if next(r.iterancestors(W_P), None) is not p:
            continue  # run belongs to a nested paragraph (text box)
        for child in r:
            if child.tag == W_T:
                segs.append([child, "t", child.text or ""])
            elif child.tag == W_TAB:
                segs.append([child, "fixed", "\t"])
            elif child.tag in (W_BR, W_CR):
                segs.append([child, "fixed", "\n"])
    return segs


def _p_text(p) -> str:
    return "".join(s[2] for s in _segments(p))


def _tokens(s: str) -> list[str]:
    return re.findall(r"\w+|\W", s)


def _apply_text(p, new: str) -> None:
    """Replace paragraph text with `new`, rewriting only the words that changed.

    A word-level diff is applied hunk by hunk (right to left), so unchanged words keep
    their exact run formatting (bold company names, italic tools, hyperlinks, tabs).
    """
    old = _p_text(p)
    if old == new:
        return
    segs = _segments(p)
    if not segs:  # empty paragraph: add a run to write into
        r = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.set(XML_SPACE, "preserve")
        r.append(t)
        p.append(r)
    a, b = _tokens(old), _tokens(new)
    a_pos = [0]
    for tok in a:
        a_pos.append(a_pos[-1] + len(tok))
    sm = difflib.SequenceMatcher(a=a, b=b, autojunk=False)
    hunks = [(a_pos[i1], a_pos[i2], "".join(b[j1:j2]))
             for op, i1, i2, j1, j2 in sm.get_opcodes() if op != "equal"]
    for start, end, text in reversed(hunks):
        _replace_range(p, start, end, text)


def _replace_range(p, start: int, end: int, middle: str) -> None:
    """Replace characters [start, end) of the paragraph text with `middle`."""
    # give every tab/break an editable text slot right before it, so text can be
    # inserted at positions that sit next to a tab (e.g. start of a tab-led line)
    for s in _segments(p):
        if s[1] == "fixed":
            prev = s[0].getprevious()
            if prev is None or prev.tag != W_T:
                slot = OxmlElement("w:t")
                slot.set(XML_SPACE, "preserve")
                s[0].addprevious(slot)
            nxt = s[0].getnext()
            if nxt is None or nxt.tag != W_T:
                slot = OxmlElement("w:t")
                slot.set(XML_SPACE, "preserve")
                s[0].addnext(slot)
    segs = _segments(p)
    text_segs = [s for s in segs if s[1] == "t"]
    middle = middle.replace("\n", " ")
    pos, spans = 0, []
    for s in segs:
        spans.append((pos, pos + len(s[2])))
        pos += len(s[2])

    # anchor: the run whose formatting the new text inherits
    anchor = None
    for s, (a, b) in zip(segs, spans):
        if s[1] == "t" and a < end and b > start and a <= start:
            anchor = s
            break
    if anchor is None:
        for s, (a, b) in zip(segs, spans):
            if s[1] == "t" and a < start <= b:   # insertion right after this run's text
                anchor = s
                break
    if anchor is None:
        for s, (a, b) in zip(segs, spans):
            if s[1] == "t" and a < end and b > start:
                anchor = s
                break
    if anchor is None:
        for s, (a, b) in zip(segs, spans):
            if s[1] == "t" and a <= start <= b:
                anchor = s
                break
    if anchor is None:
        anchor = text_segs[0]

    for s, (a, b) in zip(segs, spans):
        overlaps = a < end and b > start
        if s[1] == "fixed":
            if overlaps:
                s[0].getparent().remove(s[0])
            continue
        if not overlaps and s is not anchor:
            continue
        t = s[2]
        if s is anchor and not overlaps:
            k = min(max(0, start - a), len(t))
            left, right = t[:k], t[k:]
        else:
            left = t[:max(0, start - a)]
            right = t[max(0, end - a):] if end - a < len(t) else ""
        s[0].set(XML_SPACE, "preserve")
        if s is not anchor:
            s[0].text = left + right
            continue
        # tabs in the new text become real <w:tab/> elements
        parts = middle.split("\t")
        s[0].text = left + parts[0]
        cur = s[0]
        for part in parts[1:]:
            tab = OxmlElement("w:tab")
            cur.addnext(tab)
            nt = OxmlElement("w:t")
            nt.set(XML_SPACE, "preserve")
            nt.text = part
            tab.addnext(nt)
            cur = nt
        cur.text = (cur.text or "") + right


def _docx_kind(p, text: str) -> str:
    ppr = p.find(qn("w:pPr"))
    style = ""
    if ppr is not None:
        ps = ppr.find(qn("w:pStyle"))
        style = (ps.get(qn("w:val")) if ps is not None else "") or ""
        if ppr.find(qn("w:numPr")) is not None:
            return "bullet"
    sl = style.lower()
    if "list" in sl or BULLET_RE.match(text):
        return "bullet"
    if sl.startswith("heading") or sl == "title":
        return "heading"
    stripped = text.strip()
    letters = re.sub(r"[^A-Za-z]", "", stripped)
    if stripped and len(stripped) <= 40 and letters and letters.isupper():
        return "heading"
    return "text"


# --------------------------------------------------------------------------- Parsing
def _lines_to_paragraphs(lines: list[str]) -> list[dict]:
    """Merge wrapped lines (common in PDFs) back into logical paragraphs."""
    merged: list[str] = []
    for raw in lines:
        line = PUA_BULLET.sub("• ", raw.strip())  # Symbol-font bullets from Word PDFs
        if not line:
            continue
        prev = merged[-1] if merged else None
        prev_is_head = prev is not None and len(prev) <= 40 and re.sub(r"[^A-Za-z]", "", prev).isupper()
        is_head = len(line) <= 40 and re.sub(r"[^A-Za-z]", "", line).isupper() and any(c.isalpha() for c in line)
        is_bullet = bool(BULLET_RE.match(line))
        continues = (prev is not None and not is_head and not is_bullet and not prev_is_head and (
            (BULLET_RE.match(prev) and (line[0].islower() or not re.search(r"[.!?;:]$", prev)))
            or line[0].islower()))
        if continues:
            merged[-1] = prev + " " + line
        else:
            merged.append(line)
    paras = []
    for i, t in enumerate(merged):
        letters = re.sub(r"[^A-Za-z]", "", t)
        kind = "bullet" if BULLET_RE.match(t) else (
            "heading" if len(t) <= 40 and letters and letters.isupper() else "text")
        paras.append({"id": f"p{i}", "text": t, "kind": kind})
    return paras


CONTACT_RE = re.compile(r"[\w.+-]+@[\w-]+\.\w|\(?\d{3}\)?[\s.-]?\d{3}[\s.-]?\d{4}|linkedin|github", re.I)


def _header_first(paras: list[dict]) -> list[dict]:
    """Find the name + contact block and put it first, in reading order.

    The name is the short line just above the first contact line (email / phone / LinkedIn), not
    simply the first line: files converted from PDF often store a section title (e.g. "SUMMARY")
    before the name. Paragraph ids don't change, so export still edits the right paragraphs.
    """
    nonempty = [i for i, p in enumerate(paras) if p["text"].strip()]
    if not nonempty:
        return paras
    contact = next((i for i in nonempty[:12] if CONTACT_RE.search(paras[i]["text"]) and len(paras[i]["text"]) < 250), None)
    before = [i for i in nonempty if contact is not None and i < contact]
    name = before[-1] if before and len(paras[before[-1]]["text"].strip()) <= 60 else nonempty[0]
    end = contact if contact is not None and contact > name else name
    paras[name]["kind"] = "text"          # the name is never a section heading
    if not any(i < name for i in nonempty):
        return paras
    return paras[name:end + 1] + paras[:name] + paras[end + 1:]


def parse_resume(path: Path) -> dict:
    ext = path.suffix.lower()
    if ext == ".docx":
        try:
            doc = Document(str(path))
        except Exception as e:
            raise ValueError(f"Couldn't open the Word file (is it a real .docx and not open/locked?): {e}")
        paras = []
        for i, p in enumerate(_docx_paragraphs(doc)):
            text = _p_text(p)
            paras.append({"id": f"p{i}", "text": text, "kind": _docx_kind(p, text) if text.strip() else "empty"})
    elif ext == ".pdf":
        import pdfplumber
        lines = []
        with pdfplumber.open(str(path)) as pdf:
            for page in pdf.pages:
                lines.extend((page.extract_text(x_tolerance=1.5) or "").splitlines())
        paras = _lines_to_paragraphs(lines)
        if not paras:
            raise ValueError("No text found in the PDF — it may be a scanned image. Use a .docx or text-based PDF.")
    else:  # .txt / .md
        paras = _lines_to_paragraphs(path.read_text(encoding="utf-8", errors="replace").splitlines())

    paras = _header_first(paras)
    # annotate section (nearest heading above) for the review UI
    section = "Header"
    for p in paras:
        if p["kind"] == "heading":
            section = re.sub(r"\s*\t+\s*", " · ", p["text"].strip()).title()
        p["section"] = section
    return {"paragraphs": paras, "source_type": ext.lstrip(".")}


# --------------------------------------------------------------------------- Export
# --------------------------------------------------------------------------- File names
MAX_STEM = 80
_SEPARATORS = re.compile(r"\s*(?:[|•·\t]|\s[-–—]\s|,)\s*")
_ROLE_NOISE = re.compile(r"^(?:remote|hybrid|on-?site|in-?office|full-?time|part-?time|contract|"
                         r"(?:req|jr|r|job|id)[-#]?\d+|\d{4,}|#\d+)$", re.I)
_ROMAN = re.compile(r"^(?:i{1,3}|iv|v|vi{0,3})$", re.I)


def _ascii(s: str) -> str:
    import unicodedata
    return unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode()


def _word_case(w: str, acronyms: bool = False) -> str:
    """Title-case words typed in one case (VARRE, varre, O'BRIEN -> O'Brien); leave mixed case alone
    (McDonald, DevOps). Roman numerals stay upper (II); with acronyms, short all-caps words too (SDE, ML)."""
    if _ROMAN.match(w):
        return w.upper()
    if acronyms and w.isupper() and len(w) <= 3 and w.isalpha():
        return w
    letters = re.sub(r"[^A-Za-z]", "", w)
    if letters.isupper() or letters.islower():
        return "".join(p[:1].upper() + p[1:].lower() for p in re.split(r"([-'’])", w))
    return w


def _name_part(w: str) -> str:
    """Letters, digits and inner hyphens only: O'Brien -> OBrien, J. -> J, José -> Jose."""
    w = re.sub(r"[^A-Za-z0-9-]", "", _ascii(w)).strip("-")
    return re.sub(r"-{2,}", "-", w)


def name_words(header_line: str | None) -> list[str]:
    """The candidate's name, as file-name words, from the first line of the contact header.
    Returns [] when the line doesn't look like a name (an email, a phone number, a long sentence)."""
    if not header_line:
        return []
    first = _SEPARATORS.split(header_line.strip(), maxsplit=1)[0]        # "Jane Doe, PhD" / "Jane Doe | email"
    if re.search(r"[@\d]|https?:|www\.", first):
        return []
    words = [_name_part(_word_case(_ascii(w))) for w in first.split()]
    words = [w for w in words if w]
    return words if 1 <= len(words) <= 5 else []


def role_words(role: str | None) -> list[str]:
    """The job title as file-name words: no brackets, requisition ids, locations or work modes; max 6 words."""
    if not role or role.strip().lower() in ("not specified", "unknown", "n/a"):
        return []
    r = re.sub(r"[(\[{].*?[)\]}]", " ", role)                          # (Remote), [Hybrid], {…}
    r = re.split(r"\s[-–—|@]\s|,|\s+at\s+", r, maxsplit=1)[0]          # "- R12345", ", NYC", "at Globex"
    words = []
    for w in re.split(r"[\s/]+", r):
        if _ROLE_NOISE.match(w.strip("#.:")):
            continue
        w = re.sub(r"[^A-Za-z0-9]", "", _ascii(w.replace("++", "PP").replace("#", "Sharp")))
        if w:
            words.append(_word_case(w, acronyms=True))
    return words[:6]


def file_stem(name: list[str], role: list[str], kind: str) -> str:
    """kind: "Resume", "Cover_Letter" or "Application". Yashwanth_Varre_Software_Engineer_Resume, or just
    the kind when no name was found. Kept to MAX_STEM characters by dropping role words, then name words."""
    if not name:
        return kind
    name, role = list(name), list(role if kind != "Cover_Letter" else [])
    stem = lambda: re.sub(r"_+", "_", "_".join(name + role + [kind])).strip("_")
    while len(stem()) > MAX_STEM and role:
        role.pop()
    while len(stem()) > MAX_STEM and len(name) > 1:
        name.pop(-2 if len(name) > 2 else -1)                            # drop middle names first
    return stem()[:MAX_STEM].rstrip("_")


def output_path(out_dir: Path, stem: str, ext: str = ".docx") -> Path:
    """out_dir/stem.ext, or stem_2.ext, stem_3.ext … so an existing file is never overwritten."""
    p, n = out_dir / f"{stem}{ext}", 2
    while p.exists():
        p, n = out_dir / f"{stem}_{n}{ext}", n + 1
    return p


def output_pair(out_dir: Path, stem: str) -> tuple[Path, Path]:
    """Matching (.docx, .pdf) paths where neither file exists yet, e.g. …_Resume_2.docx / …_Resume_2.pdf."""
    cur, n = stem, 2
    while (out_dir / f"{cur}.docx").exists() or (out_dir / f"{cur}.pdf").exists():
        cur, n = f"{stem}_{n}", n + 1
    return out_dir / f"{cur}.docx", out_dir / f"{cur}.pdf"


def set_properties(doc, title: str = "", author: str = "") -> None:
    """Document title/author: Word, PDF viewers and some ATS show these. LibreOffice carries them into the PDF."""
    if title:
        doc.core_properties.title = title
    if author:
        doc.core_properties.author = author
        doc.core_properties.last_modified_by = author


def applied_text(paragraphs: list[dict], edits: dict[str, str]) -> str:
    """Plain resume text with the accepted edits swapped in."""
    return "\n".join(edits.get(p["id"], p["text"]) for p in paragraphs if edits.get(p["id"], p["text"]).strip())


def contact_header(paragraphs: list[dict], edits: dict[str, str] | None = None) -> list[str]:
    """Name + contact lines: everything above the first section heading (at most 4 lines)."""
    edits = edits or {}
    out = []
    for p in paragraphs:
        if p["kind"] == "heading":
            break
        text = edits.get(p["id"], p["text"]).replace("\t", "  ").strip()
        if text:
            out.append(text)
    return out[:4]


# --------------------------------------------------------------------------- Typography
def kw_regex(term: str) -> re.Pattern:
    """Whole words, flexible spaces/hyphens, simple plurals. Group "k" is the keyword itself."""
    t = re.escape(term.strip()).replace(r"\ ", r"[\s-]+")
    return re.compile(rf"(?:^|(?<=[^A-Za-z0-9]))(?P<k>{t}(?:s|es)?)(?=$|[^A-Za-z0-9])", re.I)


def keyword_spans(text: str, terms: list[str]) -> list[tuple[int, int]]:
    """Merged [start, end) character spans of every keyword found in text."""
    spans = sorted((m.start("k"), m.end("k")) for t in terms if t and t.strip()
                   for m in kw_regex(t).finditer(text))
    out: list[list[int]] = []
    for a, b in spans:
        if out and a <= out[-1][1]:
            out[-1][1] = max(out[-1][1], b)
        else:
            out.append([a, b])
    return [(a, b) for a, b in out]


def justify_ok(kind: str, text: str) -> bool:
    """Body paragraphs and bullets that wrap. Rows that push dates/locations right with tabs or
    separators, manual line breaks (Word stretches those lines) and short one-liners stay as they are."""
    t = text.strip()
    return (kind in ("text", "bullet") and "\n" not in t and not SEPARATOR_RE.search(t)
            and len(t) >= 90)


def body_roles(paragraphs: list[dict]) -> dict[str, str]:
    """{id: "header" | "heading" | "body"} - the name/contact block is everything above the first heading."""
    roles, seen_heading = {}, False
    for p in paragraphs:
        if p["kind"] == "heading":
            seen_heading = True
            roles[p["id"]] = "heading"
        else:
            roles[p["id"]] = "body" if seen_heading else "header"
    return roles


_RUN_TEXT = {W_T, W_TAB, W_BR, W_CR}
_RUN_SAFE = _RUN_TEXT | {qn("w:rPr"), qn("w:lastRenderedPageBreak"), qn("w:softHyphen"), qn("w:noBreakHyphen")}


def _run_text(r) -> str:
    return "".join((c.text or "") if c.tag == W_T else "\t" if c.tag == W_TAB else "\n"
                   for c in r if c.tag in _RUN_TEXT)


def _trim_run(r, lo: int, hi: int) -> None:
    """Keep only characters [lo, hi) of the run's text."""
    pos = 0
    for c in list(r):
        if c.tag not in _RUN_TEXT:
            continue
        n = len(c.text or "") if c.tag == W_T else 1
        a, b = max(lo, pos), min(hi, pos + n)
        if b <= a:
            r.remove(c)
        elif c.tag == W_T:
            c.text = (c.text or "")[a - pos:b - pos]
            c.set(XML_SPACE, "preserve")
        pos += n


def _direct_runs(p) -> list:
    return [r for r in p.iter(W_R) if next(r.iterancestors(W_P), None) is p]


def _set_bold(r) -> None:
    rpr = r.get_or_add_rPr()   # python-docx inserts children in schema order, which Word insists on
    for el in (rpr.get_or_add_b(), rpr.get_or_add_bCs()):
        el.attrib.pop(qn("w:val"), None)


def _keep_whole(r) -> None:
    """Never hyphenate this run: "no proofing" plus no language, which Word and LibreOffice both skip when
    hyphenating. Keywords then stay whole on the line, so a PDF parser never reads "micro-services"."""
    rpr = r.get_or_add_rPr()
    rpr.get_or_add_noProof()
    lang = rpr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        later = next((rpr.find(qn(t)) for t in ("w:eastAsianLayout", "w:specVanish", "w:oMath")
                      if rpr.find(qn(t)) is not None), None)
        later.addprevious(lang) if later is not None else rpr.append(lang)
    lang.set(qn("w:val"), "zxx")


def bold_keywords(p, terms: list[str], bold: bool = True) -> int:
    """Make only the keyword characters bold (when bold), splitting runs at their edges, and keep every
    keyword from being hyphenated. Surrounding text keeps its formatting; runs holding anything but text
    (images, fields) are left alone. Returns the number of keywords found."""
    runs = _direct_runs(p)
    text = "".join(_run_text(r) for r in runs)
    spans = keyword_spans(text, terms)
    if not spans:
        return 0
    cuts = sorted({x for span in spans for x in span})
    pos = 0
    for r in runs:
        n = len(_run_text(r))
        inner = [c - pos for c in cuts if pos < c < pos + n]
        if inner and all(c.tag in _RUN_SAFE for c in r):
            pieces, prev = [], 0
            for k in inner + [n]:
                piece = copy.deepcopy(r)
                _trim_run(piece, prev, k)
                pieces.append(piece)
                prev = k
            for piece in reversed(pieces):
                r.addnext(piece)
            r.getparent().remove(r)
        pos += n
    pos, done = 0, 0
    for r in _direct_runs(p):
        n = len(_run_text(r))
        if n and any(a <= pos and pos + n <= b for a, b in spans) and all(c.tag in _RUN_SAFE for c in r):
            if bold:
                _set_bold(r)
            _keep_whole(r)
        pos += n
    return len(spans)


def _justify(p) -> None:
    jc = p.get_or_add_pPr().get_or_add_jc()
    if jc.get(qn("w:val")) not in (None, "left", "start", "both"):
        return   # centred / right-aligned on purpose
    jc.set(qn("w:val"), "both")


def _font_attrs(rfonts, font: str) -> None:
    for k in ("asciiTheme", "hAnsiTheme", "cstheme", "eastAsiaTheme"):
        rfonts.attrib.pop(qn(f"w:{k}"), None)
    for k in ("ascii", "hAnsi", "cs", "eastAsia"):
        rfonts.set(qn(f"w:{k}"), font)


def _ensure_rfonts(rpr, font: str) -> None:
    _font_attrs(rpr.get_or_add_rFonts(), font)


def apply_font(doc, font: str) -> None:
    """One font everywhere: document defaults, every style and every run (theme fonts removed, so Word
    doesn't fall back to Calibri/Aptos). Sizes, bold, italic and colours are untouched."""
    styles = doc.styles.element
    for rf in styles.iter(qn("w:rFonts")):
        _font_attrs(rf, font)
    defaults = styles.find(qn("w:docDefaults"))
    if defaults is None:
        defaults = OxmlElement("w:docDefaults")
        styles.insert(0, defaults)
    rpd = defaults.find(qn("w:rPrDefault"))
    if rpd is None:
        rpd = OxmlElement("w:rPrDefault")
        defaults.insert(0, rpd)
    rpr = rpd.find(qn("w:rPr"))
    if rpr is None:
        rpr = OxmlElement("w:rPr")
        rpd.append(rpr)
    _ensure_rfonts(rpr, font)
    _ensure_rfonts(doc.styles["Normal"].element.get_or_add_rPr(), font)
    for r in doc.element.body.iter(W_R):
        _ensure_rfonts(r.get_or_add_rPr(), font)


_SETTINGS_AFTER_HYPHENATION = ("w:showEnvelope", "w:summaryLength", "w:clickAndTypeStyle", "w:defaultTableStyle",
                               "w:evenAndOddHeaders", "w:drawingGridHorizontalSpacing", "w:drawingGridVerticalSpacing",
                               "w:displayHorizontalDrawingGridEvery", "w:displayVerticalDrawingGridEvery",
                               "w:doNotUseMarginsForDrawingGridOrigin", "w:doNotShadeFormData", "w:noPunctuationKerning",
                               "w:characterSpacingControl", "w:footnotePr", "w:endnotePr", "w:compat", "w:docVars",
                               "w:rsids", "m:mathPr", "w:themeFontLang", "w:clrSchemeMapping", "w:decimalSymbol", "w:listSeparator")


def hyphenate(doc) -> None:
    """Turn on Word's automatic hyphenation so justified lines don't open wide gaps. Display only: the text
    an ATS reads is unchanged. ALL-CAPS words (AWS, SQL) are never split; at most 2 hyphens in a row."""
    st = doc.settings.element
    for tag in ("w:autoHyphenation", "w:consecutiveHyphenLimit", "w:hyphenationZone", "w:doNotHyphenateCaps"):
        for el in st.findall(qn(tag)):
            st.remove(el)
    els = [OxmlElement("w:autoHyphenation"), OxmlElement("w:consecutiveHyphenLimit"), OxmlElement("w:doNotHyphenateCaps")]
    els[1].set(qn("w:val"), "2")
    anchor = st.find(qn("w:defaultTabStop"))
    if anchor is not None:
        for el in reversed(els):
            anchor.addnext(el)
        return
    later = next((st.find(qn(t)) for t in _SETTINGS_AFTER_HYPHENATION if st.find(qn(t)) is not None), None)
    for el in els:
        later.addprevious(el) if later is not None else st.append(el)


def finish_resume(doc, paragraphs: list[dict], elements: dict, keywords: list[str],
                  font: str = DOC_FONT, bold: bool = True) -> None:
    """elements: {paragraph_id: w:p}. Justify body text, bold keywords, set the font."""
    roles, justified = body_roles(paragraphs), False
    for p in paragraphs:
        el = elements.get(p["id"])
        if el is None or roles[p["id"]] != "body":
            continue
        text = _p_text(el)
        if justify_ok(p["kind"], text):
            _justify(el)
            justified = True
        if keywords:
            bold_keywords(el, keywords, bold)
    if justified:
        hyphenate(doc)
    if font and font != KEEP_FONT:
        apply_font(doc, font)


def export_docx(source: Path, source_type: str, paragraphs: list[dict],
                edits: dict[str, str], dest: Path, keywords: list[str] | None = None,
                font: str = DOC_FONT, bold: bool = True, title: str = "", author: str = "") -> Path:
    """edits: {paragraph_id: new_text} for the accepted changes only. keywords: job-posting terms
    (and variants) to show in bold."""
    dest.parent.mkdir(parents=True, exist_ok=True)
    final = [{**p, "text": edits.get(p["id"], p["text"])} for p in paragraphs]
    if source_type == "docx":
        tmp = dest.with_suffix(".tmp.docx")
        shutil.copyfile(source, tmp)
        doc = Document(str(tmp))
        ps = _docx_paragraphs(doc)
        for pid, new in edits.items():
            idx = int(pid[1:])
            if 0 <= idx < len(ps):
                _apply_text(ps[idx], new)
        elements = {p["id"]: ps[int(p["id"][1:])] for p in final if int(p["id"][1:]) < len(ps)}
        finish_resume(doc, final, elements, keywords or [], font, bold)
        set_properties(doc, title, author)
        doc.save(str(dest))
        tmp.unlink(missing_ok=True)
        return dest
    return _build_clean_docx(final, dest, keywords or [], font, bold, title, author)


def _bottom_border(paragraph) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    for k, v in (("val", "single"), ("sz", "6"), ("space", "1"), ("color", "808080")):
        b.set(qn(f"w:{k}"), v)
    bdr.append(b)
    ppr.append(bdr)


def _build_clean_docx(paras: list[dict], dest: Path, keywords: list[str] | None = None,
                      font: str = DOC_FONT, bold: bool = True, title: str = "", author: str = "") -> Path:
    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Inches(0.6)
        s.left_margin = s.right_margin = Inches(0.7)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri" if font == KEEP_FONT else font
    normal.font.size = Pt(10.5 if font == KEEP_FONT else 11.5)   # EB Garamond runs small
    normal.paragraph_format.space_after = Pt(2)

    elements = {}
    first_heading = next((i for i, p in enumerate(paras) if p["kind"] == "heading"), len(paras))
    for i, p in enumerate(paras):
        text = p["text"].strip()
        if not text:
            continue
        if i == 0:  # name
            para = doc.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = para.add_run(text)
            r.bold = True
            r.font.size = Pt(18)
        elif i < first_heading:  # contact lines
            para = doc.add_paragraph(text)
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif p["kind"] == "heading":
            para = doc.add_paragraph()
            para.paragraph_format.space_before = Pt(10)
            para.paragraph_format.space_after = Pt(4)
            r = para.add_run(text.upper())
            r.bold = True
            r.font.size = Pt(11.5)
            r.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
            _bottom_border(para)
        elif p["kind"] == "bullet":
            para = doc.add_paragraph(BULLET_RE.sub("", text), style="List Bullet")
            para.paragraph_format.space_after = Pt(1)
        else:
            para = doc.add_paragraph(text)
        elements[p["id"]] = para._p
    finish_resume(doc, paras, elements, keywords or [], font, bold)
    set_properties(doc, title, author)
    doc.save(str(dest))
    return dest


def export_cover_letter(letter: dict, header: list[str], dest: Path, keywords: list[str] | None = None,
                        font: str = DOC_FONT, bold: bool = False, title: str = "", author: str = "") -> Path:
    """Write the cover letter as a simple one-page business letter. Body paragraphs are justified;
    keywords are bolded only when asked (off by default for letters)."""
    from datetime import date
    dest.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    for sec in doc.sections:
        sec.top_margin = sec.bottom_margin = Inches(0.9)
        sec.left_margin = sec.right_margin = Inches(1)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri" if font == KEEP_FONT else font
    normal.font.size = Pt(11 if font == KEEP_FONT else 12)
    normal.paragraph_format.space_after = Pt(10)
    normal.paragraph_format.line_spacing = 1.1

    if header:
        para = doc.add_paragraph()
        para.paragraph_format.space_after = Pt(0)
        r = para.add_run(header[0])
        r.bold = True
        r.font.size = Pt(16)
        for line in header[1:]:
            q = doc.add_paragraph(line)
            q.paragraph_format.space_after = Pt(0)
            q.runs[0].font.size = Pt(10)
            q.runs[0].font.color.rgb = RGBColor(0x55, 0x55, 0x55)
        _bottom_border(doc.paragraphs[-1])
        doc.paragraphs[-1].paragraph_format.space_after = Pt(14)
    today = date.today()
    doc.add_paragraph(f"{today:%B} {today.day}, {today.year}")
    doc.add_paragraph(letter.get("greeting", "").strip() or "Dear Hiring Team,")
    for body in letter.get("paragraphs", []):
        if body.strip():
            para = doc.add_paragraph(body.strip())
            _justify(para._p)
            if keywords:
                bold_keywords(para._p, keywords, bold)
    closing = doc.add_paragraph(letter.get("closing", "").strip() or "Sincerely,")
    closing.paragraph_format.space_before = Pt(4)
    closing.paragraph_format.space_after = Pt(2)
    doc.add_paragraph(letter.get("signature", "").strip() or (header[0] if header else ""))
    hyphenate(doc)
    if font != KEEP_FONT:
        apply_font(doc, font)
    set_properties(doc, title, author)
    doc.save(str(dest))
    return dest
